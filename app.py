import streamlit as st
import asyncio
import subprocess
import sys
import requests
import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

from scraper    import scrape_portfolio, get_suggestions, DATABASE, JS_VISIBLE_TEXT, extract_with_groq, find_in_database
from sebi_eve   import search_sebi, get_pdf_url, download_pdf, extract_sections

# ── Install Playwright browser once ──────────────────────────
@st.cache_resource
def install_playwright_browser():
    result = subprocess.run(
        [sys.executable, "-m", "playwright", "install", "chromium"],
        capture_output=True, text=True
    )
    return result.returncode

install_playwright_browser()

# ── Page config ───────────────────────────────────────────────
st.set_page_config(
    page_title="ASK Eve",
    page_icon="🔮",
    layout="centered"
)

# ── Global CSS ────────────────────────────────────────────────
st.markdown("""
<style>
    .main { max-width: 800px; margin: auto; }

    /* Header */
    .ask-header {
        background-color: #033A49;
        padding: 20px 28px;
        border-radius: 10px;
        margin-bottom: 28px;
        display: flex;
        align-items: center;
        justify-content: space-between;
    }
    .ask-header-left  { display: flex; align-items: center; gap: 16px; }
    .ask-gold-bar     { width: 4px; min-height: 46px; background-color: #BB962C; border-radius: 2px; flex-shrink: 0; }
    .ask-header-title { font-size: 22px; font-weight: 600; color: white; margin: 0 0 3px; }
    .ask-header-sub   { font-size: 12px; color: rgba(255,255,255,0.55); margin: 0; }
    .ask-badge        { background-color: #BB962C; color: white; font-size: 12px; font-weight: 600; padding: 5px 12px; border-radius: 20px; white-space: nowrap; }

    /* Buttons */
    .stButton > button {
        background-color: #033A49;
        color: white;
        border: none;
        padding: 0.55rem 1.4rem;
        font-size: 14px;
        font-weight: 500;
        border-radius: 8px;
        margin-top: 4px;
        width: 100%;
    }
    .stButton > button:hover { background-color: #044f65; color: white; border: none; }

    /* Footer */
    .ask-footer { text-align: center; color: #aaa; font-size: 12px; margin-top: 40px; padding-top: 14px; border-top: 1px solid #f0f0f0; }

    /* Eve specific */
    .eve-result-card {
        border: 0.5px solid #e0e0e0;
        border-radius: 8px;
        padding: 10px 16px;
        margin-bottom: 8px;
        cursor: pointer;
        transition: background 0.15s;
    }
    .eve-result-card:hover { background: #f0f7f9; }
    .eve-section-log { font-size: 13px; line-height: 1.8; }
</style>
""", unsafe_allow_html=True)

# ── Session state ─────────────────────────────────────────────
for key, default in [
    ("search_count", 0),
    ("firm_to_search", ""),
    ("eve_results", []),
    ("eve_selected", None),
    ("eve_query", ""),
]:
    if key not in st.session_state:
        st.session_state[key] = default

# ── Header ────────────────────────────────────────────────────
st.markdown(f"""
<div class="ask-header">
    <div class="ask-header-left">
        <div class="ask-gold-bar"></div>
        <div>
            <p class="ask-header-title">ASK Eve</p>
            <p class="ask-header-sub">Alternative Investments Intelligence Platform</p>
        </div>
    </div>
    <div class="ask-badge">🔍 {st.session_state.search_count} searches</div>
</div>
""", unsafe_allow_html=True)

# ── API key ───────────────────────────────────────────────────
try:
    groq_key = st.secrets["GROQ_API_KEY"]
except Exception:
    st.error('Groq API key missing. Add to Streamlit Secrets:\nGROQ_API_KEY = "gsk_..."')
    st.stop()

# ── Two tabs ──────────────────────────────────────────────────
tab1, tab2 = st.tabs(["📊  Portscope", "📄  DRHP Extractor"])


# ════════════════════════════════════════════════════════════
#  TAB 1 — PORTSCOPE
# ════════════════════════════════════════════════════════════
with tab1:

    with st.expander("🔧 Debug info"):
        col1, col2 = st.columns(2)
        with col1:
            if st.button("Test Groq API", key="t_groq"):
                try:
                    r = requests.post(
                        "https://api.groq.com/openai/v1/chat/completions",
                        headers={"Authorization": f"Bearer {groq_key}", "Content-Type": "application/json"},
                        json={"model": "llama-3.1-8b-instant",
                              "messages": [{"role": "user", "content": "Say: working"}],
                              "max_tokens": 10},
                        timeout=15
                    )
                    r.raise_for_status()
                    st.success(f"Groq OK: {r.json()['choices'][0]['message']['content'][:40]}")
                except Exception as e:
                    st.error(f"Groq failed: {e}")

        with col2:
            if st.button("Test Browser", key="t_browser"):
                async def _test():
                    from playwright.async_api import async_playwright
                    async with async_playwright() as p:
                        b = await p.chromium.launch(headless=True, args=["--no-sandbox","--disable-dev-shm-usage"])
                        pg = await b.new_page()
                        await pg.goto("https://example.com", timeout=15000)
                        t = await pg.title()
                        await b.close()
                        return t
                try:
                    loop = asyncio.new_event_loop(); asyncio.set_event_loop(loop)
                    t = loop.run_until_complete(_test()); loop.close()
                    st.success(f"Browser OK: {t}")
                except Exception as e:
                    st.error(f"Browser failed: {e}")

        st.markdown("---")
        st.markdown("**Step-by-step debugger**")
        dbg = st.text_input("Firm name to debug:", key="dbg_input", placeholder="e.g. KKR, EQT...")
        if st.button("Run debug", key="run_dbg") and dbg.strip():
            firm = dbg.strip()
            db = find_in_database(firm)
            if db:
                st.success(f"✅ In database — URL: `{db['url']}` | Pagination: `{db.get('pagination')}`")
                portfolio_url = db["url"]
            else:
                st.warning("⚠️ Not in database")
                portfolio_url = None

            if portfolio_url:
                async def _debug_fetch(url):
                    from playwright.async_api import async_playwright
                    async with async_playwright() as p:
                        b = await p.chromium.launch(headless=True, args=["--no-sandbox","--disable-dev-shm-usage","--disable-gpu"])
                        ctx = await b.new_context(
                            user_agent="Mozilla/5.0 Chrome/120.0.0.0 Safari/537.36",
                            viewport={"width": 1280, "height": 900}
                        )
                        pg = await ctx.new_page()
                        await pg.route("**/*.{png,jpg,jpeg,gif,webp,svg,woff,woff2,ttf}", lambda r: r.abort())
                        await pg.goto(url, wait_until="domcontentloaded", timeout=45000)
                        await pg.wait_for_timeout(4000)
                        title = await pg.title()
                        await pg.evaluate("""() => {
                            ['[class*="cookie"]','[class*="consent"]'].forEach(s=>
                                document.querySelectorAll(s).forEach(e=>e.remove()));
                        }""")
                        for _ in range(5):
                            await pg.evaluate("window.scrollTo(0,document.body.scrollHeight)")
                            await pg.wait_for_timeout(1000)
                        vis = await pg.evaluate(JS_VISIBLE_TEXT)
                        await ctx.close(); await b.close()
                        return title, vis
                try:
                    loop = asyncio.new_event_loop(); asyncio.set_event_loop(loop)
                    title, vis = loop.run_until_complete(_debug_fetch(portfolio_url)); loop.close()
                    st.success(f"✅ Page loaded: {title} — {len(vis)} visible text items")
                    st.code("\n".join(vis[:60]))
                    st.markdown("**Groq extraction:**")
                    found = extract_with_groq("\n".join(vis[:300]), firm, groq_key)
                    if found:
                        st.success(f"✅ {len(found)} companies found")
                        for c in found: st.write(f"  • {c}")
                    else:
                        st.error("❌ Groq returned nothing")
                except Exception as e:
                    st.error(f"Debug error: {e}")

    # Search
    firm_input = st.text_input(
        "Investor / PE firm name",
        placeholder="Type name — e.g. EQT, Actis, Carlyle...",
        key="firm_input"
    )

    if firm_input:
        sugs = get_suggestions(firm_input)
        if sugs:
            st.markdown("**Suggestions from database:**")
            cols = st.columns(min(len(sugs), 4))
            for i, s in enumerate(sugs[:8]):
                with cols[i % 4]:
                    if st.button(s, key=f"sug_{s}"):
                        st.session_state.firm_to_search = s
                        st.rerun()
            st.caption("Click a suggestion, or type the full name and click Search.")

    if st.button("🔍 Search Portfolio", key="search_btn"):
        if firm_input.strip():
            st.session_state.firm_to_search = firm_input.strip()

    fts = st.session_state.firm_to_search
    if fts and (st.session_state.get("search_btn") or fts.lower() != firm_input.strip().lower()):
        st.session_state.search_count += 1
        status = st.empty()
        companies = []
        try:
            loop = asyncio.new_event_loop(); asyncio.set_event_loop(loop)
            companies = loop.run_until_complete(scrape_portfolio(fts, groq_key, status))
            loop.close()
        except Exception as e:
            st.error(f"Error: {e}")
        status.empty()

        # Check bot blocked
        if len(companies) == 1 and companies[0].startswith("BOT_BLOCKED:"):
            blocked_url = companies[0].replace("BOT_BLOCKED:", "")
            st.warning(f"⚠️ {fts}'s website uses bot protection — automated access is blocked.")
            st.markdown(f"**Visit their portfolio page directly:**\n👉 [{blocked_url}]({blocked_url})")
        elif companies:
            st.success(f"✅ Found {len(companies)} companies in {fts}'s portfolio")
            st.markdown("---")

            # Excel
            wb = openpyxl.Workbook(); ws = wb.active
            ws.title = f"{fts} Portfolio"
            ws["C4"] = fts.upper(); ws["C4"].font = Font(bold=True, size=12)
            ws["C5"] = "SR NO"; ws["D5"] = "COMPANY"
            ws["C5"].font = Font(bold=True, color="FFFFFF"); ws["C5"].fill = PatternFill("solid", fgColor="033A49"); ws["C5"].alignment = Alignment(horizontal="center")
            ws["D5"].font = Font(bold=True, color="FFFFFF"); ws["D5"].fill = PatternFill("solid", fgColor="033A49")
            for i, c in enumerate(companies, 1):
                ws[f"C{5+i}"] = i; ws[f"D{5+i}"] = c
                ws[f"C{5+i}"].alignment = Alignment(horizontal="center")
            ws.column_dimensions["C"].width = 10; ws.column_dimensions["D"].width = 45
            buf = io.BytesIO(); wb.save(buf); buf.seek(0)

            c1, c2 = st.columns(2)
            with c1:
                st.download_button("⬇️ Download Excel", buf.getvalue(),
                    f"{fts.replace(' ','_')}_portfolio.xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            with c2:
                st.download_button("⬇️ Download CSV",
                    "SR NO,COMPANY\n" + "\n".join(f"{i+1},{c}" for i,c in enumerate(companies)),
                    f"{fts.replace(' ','_')}_portfolio.csv", "text/csv")

            st.markdown(f"### {fts.upper()} — Portfolio Companies")
            rows = "".join(
                f'<tr style="background:{"#f8f8f8" if i%2==0 else "#fff"}">'
                f'<td style="padding:8px 14px;border:1px solid #ddd;text-align:center;color:#555;width:70px">{i}</td>'
                f'<td style="padding:8px 14px;border:1px solid #ddd;color:#222">{c}</td></tr>'
                for i, c in enumerate(companies, 1)
            )
            st.markdown(f"""
            <table style="width:100%;border-collapse:collapse;font-size:14px;margin-top:8px">
                <thead><tr style="background:#033A49">
                    <th style="padding:10px 14px;border:1px solid #ddd;color:white;text-align:center">SR NO</th>
                    <th style="padding:10px 14px;border:1px solid #ddd;color:white;text-align:left">COMPANY</th>
                </tr></thead>
                <tbody>{rows}</tbody>
            </table>""", unsafe_allow_html=True)
        else:
            st.error(f"No companies found for '{fts}'.")
            st.markdown("Open 🔧 Debug info → type firm name → click **Run debug**")

        st.session_state.firm_to_search = ""


# ════════════════════════════════════════════════════════════
#  TAB 2 — ASK Eve DRHP Extractor
# ════════════════════════════════════════════════════════════
with tab2:
    st.markdown("#### DRHP Section Extractor")
    st.markdown(
        "Search for any company that has filed a DRHP with SEBI. "
        "ASK Eve will extract the key sections for you — no need to read 400 pages."
    )

    # Sections info
    with st.expander("📋 Which sections are extracted?"):
        st.markdown("""
        | Section | What it covers | Max pages |
        |---|---|---|
        | **Company Overview / Our Business** | What the company does, business model, operations | 25 |
        | **Risk Factors** | Key risks the company has disclosed | 20 |
        | **Promoters & Promoter Group** | Who the promoters are, their background | 20 |
        | **Objects of the Offer** | What the IPO money will be used for | 10 |
        | **Financial Information** | Restated financial statements | 35 |
        """)

    st.markdown("---")

    # Search input
    eve_query = st.text_input(
        "Company name",
        placeholder="e.g. Mangalam Drugs, Playsimple Games, PhonePe...",
        key="eve_query_input",
        value=st.session_state.eve_query
    )

    if st.button("🔍 Search SEBI Filings", key="eve_search"):
        if not eve_query.strip():
            st.warning("Please enter a company name.")
        else:
            with st.spinner(f"Searching SEBI for '{eve_query}'..."):
                try:
                    loop = asyncio.new_event_loop(); asyncio.set_event_loop(loop)
                    results = loop.run_until_complete(search_sebi(eve_query.strip()))
                    loop.close()
                    st.session_state.eve_results  = results
                    st.session_state.eve_selected = None
                    st.session_state.eve_query    = eve_query.strip()
                except Exception as e:
                    st.error(f"Search error: {e}")
                    results = []

            if not results:
                st.error(f"No DRHP filings found for '{eve_query}'. Try a different name or check spelling.")

    # Show search results
    if st.session_state.eve_results:
        results = st.session_state.eve_results
        st.markdown(f"**{len(results)} filing(s) found — click to select:**")

        for i, r in enumerate(results):
            btn_label = f"📄  {r['name']}  ·  {r['date']}"
            if st.button(btn_label, key=f"eve_result_{i}"):
                st.session_state.eve_selected = r
                st.session_state.eve_query    = r["name"]
                st.rerun()

    # Show selected company + extract button
    if st.session_state.eve_selected:
        sel = st.session_state.eve_selected
        st.markdown("---")
        st.success(f"✅ Selected: **{sel['name']}** — filed {sel['date']}")
        st.caption(f"SEBI page: {sel['url']}")

        if st.button("⚡ Extract DRHP Sections", key="eve_extract"):
            st.session_state.search_count += 1
            log = []

            try:
                # Step 1: Get PDF URL
                with st.spinner("Step 1/3 — Finding PDF on SEBI..."):
                    loop = asyncio.new_event_loop(); asyncio.set_event_loop(loop)
                    pdf_url = loop.run_until_complete(get_pdf_url(sel["url"]))
                    loop.close()

                if not pdf_url:
                    st.error("Could not find the PDF link on SEBI. The page structure may have changed.")
                    st.stop()

                st.info(f"📎 PDF found: `{pdf_url[:80]}...`")

                # Step 2: Download PDF
                with st.spinner("Step 2/3 — Downloading DRHP PDF..."):
                    pdf_bytes = download_pdf(pdf_url)
                    size_mb = len(pdf_bytes) / (1024 * 1024)
                    st.info(f"📥 Downloaded: {size_mb:.1f} MB")

                # Step 3: Extract sections
                with st.spinner("Step 3/3 — Extracting target sections..."):
                    output_pdf, log = extract_sections(pdf_bytes)

                # Results
                st.success(f"✅ Extraction complete! Output PDF ready.")
                st.markdown("**Extraction log:**")
                for line in log:
                    st.markdown(f"<div class='eve-section-log'>{line}</div>", unsafe_allow_html=True)

                st.markdown("---")
                st.download_button(
                    label="⬇️ Download Extracted PDF",
                    data=output_pdf,
                    file_name=f"{sel['name'].replace(' ', '_')}_ASKEve_Extract.pdf",
                    mime="application/pdf",
                    type="primary"
                )
                st.caption(
                    "This PDF contains only the extracted sections from the original DRHP. "
                    "Pages are exactly as filed with SEBI."
                )

            except Exception as e:
                st.error(f"Extraction failed: {e}")
                st.markdown("Try searching again or check if the SEBI page is accessible.")

# ── Footer ────────────────────────────────────────────────────
st.markdown(
    "<div class='ask-footer'>ASK Eve &nbsp;·&nbsp; Built by Aaryaman Singh &nbsp;·&nbsp; "
    "Data sourced from SEBI</div>",
    unsafe_allow_html=True
)
